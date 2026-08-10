# -*- coding: utf-8 -*-
"""Rebuild master.xlsx from the Word/Markdown export of MatjarLink_Master_Data_File.

The export slices every wide sheet vertically into separate markdown tables. Each slice
keeps the sheet's row order AND row count (header rows appear as all-empty rows in the
right-hand slices), so re-joining is positional -- slice A row i belongs with slice B
row i. Row counts are asserted before joining; nothing is inferred.

Usage:  python3 rebuild_master.py <export.md> [master.xlsx]
"""
import re
import sys

import openpyxl

SEP = re.compile(r'^\|[\s:|-]*-[\s:|-]*\|$')   # alignment row: must contain a dash


def clean(c):
    c = c.replace('\\_', '_').replace('\\*', '*').replace('\\.', '.')
    c = re.sub(r'\*\*(.*?)\*\*', r'\1', c)
    return c.replace('﻿', '').strip()


def split_cells(s):
    """Split a markdown row into cells, repairing labels that embed a literal '|'.

    Category and group labels are written as "English  |  عربي" and the exporter did
    not escape that pipe, so those cells arrive split in two. Both halves are wrapped
    by a single bold span, so a cell whose '**' markers are unbalanced continues into
    the next cell -- that is the join signal, and it is exact rather than heuristic.
    """
    raw = s.split('|')[1:-1]
    merged, buf = [], None
    for c in raw:
        cur = c if buf is None else buf + '|' + c
        if cur.count('**') % 2:          # unbalanced -> label continues
            buf = cur
            continue
        buf = None
        merged.append(cur)
    if buf is not None:                  # never closed; keep as-is
        merged.append(buf)
    return [clean(c) for c in merged]


ARABIC = re.compile(r'^[؀-ۿ]')
FLAG = ('Yes', 'No', '')
DTYPE = ('Dropdown', 'Boolean', 'Decimal', 'Integer', 'Text', 'List', 'Free', '')


def fix_label(cells, width):
    """Left slices: a label written "English  |  عربي" without bold markers arrives
    split. The continuation half always starts with Arabic script -- merge on that."""
    while len(cells) > width:
        for i in range(len(cells) - 1):
            if cells[i] and ARABIC.match(cells[i + 1]):
                cells = cells[:i] + [f'{cells[i]}  |  {cells[i + 1]}'] + cells[i + 2:]
                break
        else:
            return cells
    return cells


def fix_list(cells, width, ntrail):
    """Right slices: an option value containing a literal '|' (e.g. "Xbox Series X|S")
    arrives split. Data_Type leads and Yes/No flags trail, so everything between those
    anchors is one list cell -- rejoin it with the pipe it lost."""
    while len(cells) > width:
        lead = 1 if cells and cells[0] in DTYPE and width > ntrail + 1 else 0
        if all(c in FLAG for c in cells[len(cells) - ntrail:]):
            i = lead
            cells = cells[:i] + ['|'.join(cells[i:len(cells) - ntrail])] + cells[len(cells) - ntrail:]
        else:
            return cells
    return cells


def slice_rows(lines, a, b, width=None, ntrail=None):
    """Table rows of lines[a:b], separators dropped, blank rows KEPT (they carry order)."""
    out = []
    for ln in lines[a:b]:
        s = ln.rstrip('\n')
        if not s.startswith('|') or SEP.match(s):
            continue
        cells = split_cells(s)
        if width and len(cells) > width:
            cells = fix_list(cells, width, ntrail) if ntrail else fix_label(cells, width)
        out.append(cells)
    return out


def join(*slices):
    """Positional column-wise join. All slices must have equal row counts."""
    n = {len(s) for s in slices}
    assert len(n) == 1, f'row-count mismatch across slices: {[len(s) for s in slices]}'
    out = []
    for parts in zip(*slices):
        row = []
        for p in parts:
            row.extend(p)
        out.append(row)
    return out


def pad(rows, width):
    return [r + [''] * (width - len(r)) if len(r) < width else r[:width] for r in rows]


def main():
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else 'master.xlsx'
    lines = open(src, encoding='utf-8').readlines()

    def hdr(pattern):
        rx = re.compile(pattern)
        for i, ln in enumerate(lines):
            if rx.match(clean(ln.rstrip('\n'))):
                return i
        raise SystemExit(f'header not found: {pattern}')

    # ---- Categories: Code|L1|L2  +  Leaf(L3)|Level
    c1 = hdr(r'^\|Code\|Top Category \(L1\)\|Sub-Category \(L2\)\|$')
    c2 = hdr(r'^\|Leaf Category \(L3\)\|Level\|$')
    c3 = hdr(r'^\|Brand\|$')
    cats = join(pad(slice_rows(lines, c1 + 1, c2, 3), 3),
                pad(slice_rows(lines, c2 + 1, c3, 2), 2))

    # ---- Product Properties: Category|Group|EN|AR  +  Data_Type|Item_List|Required
    p1 = hdr(r'^\|Category\|Property Group\|Property \(EN\)\|Property \(AR\)\|$')
    p2 = hdr(r'^\|Data_Type\|Item_List\|Required\|$')
    p3 = hdr(r'^\|Category\|Property Group\|Property \(EN\)\|Property \(AR\)\|Data\|$')
    pp = join(pad(slice_rows(lines, p1 + 1, p2, 4), 4),
              pad(slice_rows(lines, p2 + 1, p3, 3, ntrail=1), 3))

    # ---- Properties + Options: Category|Group|EN|AR|Data + Values|Required|Option
    q1 = p3
    q2 = hdr(r'^\|Values \(Item List\)\|Required\|Option\|$')
    q3 = hdr(r'^\|Activity Code\|Class\|Section\|Division\|$')
    po = join(pad(slice_rows(lines, q1 + 1, q2, 5), 5),
              pad(slice_rows(lines, q2 + 1, q3, 3, ntrail=2), 3))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet('Categories')
    ws.append(['Code', 'Top Category (L1)', 'Sub-Category (L2)',
               'Leaf Category (L3)', 'Level'])
    for r in cats:
        ws.append(r)

    ws = wb.create_sheet('Product Properties')
    ws.append(['Category', 'Property Group', 'Property (EN)', 'Property (AR)',
               'Data_Type', 'Item_List', 'Required'])
    for r in pp:
        ws.append(r)

    ws = wb.create_sheet('Properties + Options')
    ws.append(['Category', 'Property Group', 'Property (EN)', 'Property (AR)',
               'Data', 'Values (Item List)', 'Required', 'Option'])
    for r in po:
        ws.append(r)

    wb.save(out)
    print(f'wrote {out}')
    print(f'  Categories           {len(cats):6d} rows  '
          f'({sum(1 for r in cats if "Leaf" in str(r[4]))} leaves)')
    print(f'  Product Properties   {len(pp):6d} rows  '
          f'({sum(1 for r in pp if r[2])} properties, '
          f'{sum(1 for r in pp if r[0])} categories)')
    print(f'  Properties + Options {len(po):6d} rows  '
          f'({sum(1 for r in po if r[2])} properties, '
          f'{sum(1 for r in po if r[0])} categories)')


if __name__ == '__main__':
    main()
